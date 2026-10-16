from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "docs" / "assets" / "manual-trabajador"
FLOW = ASSETS / "flujo"
WIDE = ASSETS / "amplio"
OUT = ROOT / "docs" / "Manual_de_uso_ETRAL_para_trabajadores_corregido.docx"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(11), Inches(8.5)
sec.left_margin = sec.right_margin = Inches(.55)
sec.top_margin, sec.bottom_margin = Inches(.48), Inches(.5)
sec.header_distance, sec.footer_distance = Inches(.2), Inches(.22)

normal = doc.styles["Normal"]
normal.font.name, normal.font.size = "Aptos", Pt(9.5)
normal.font.color.rgb = RGBColor(30, 41, 59)
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.line_spacing = 1.02
for name, size in [("Title", 25), ("Heading 1", 17), ("Heading 2", 12)]:
    style = doc.styles[name]
    style.font.name, style.font.size, style.font.bold = "Aptos Display", Pt(size), True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(5)

header = sec.header.paragraphs[0]
header.text = "ETRAL  |  MANUAL DE USO POR FLUJO"
header.runs[0].font.size = Pt(8)
header.runs[0].font.bold = True
header.runs[0].font.color.rgb = RGBColor(0, 0, 0)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run("Manual del sistema ETRAL  |  ")
run.font.size = Pt(8)
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
footer._p.append(fld)

def text(value="", style=None):
    return doc.add_paragraph(value, style)

def numbered(number, value):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(.23)
    p.paragraph_format.first_line_indent = Inches(-.23)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(f"{number}. ")
    r.bold = True
    r.font.size = Pt(8.5)
    r = p.add_run(value)
    r.font.size = Pt(8.5)

def add_legend(items):
    if len(items) <= 4:
        for number, value in enumerate(items, 1):
            numbered(number, value)
        return
    split = (len(items) + 1) // 2
    rows = max(split, len(items) - split)
    table = doc.add_table(rows=rows, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(4.7)
    table.columns[1].width = Inches(4.7)
    for col, group in enumerate((items[:split], items[split:])):
        for row, value in enumerate(group):
            number = row + 1 if col == 0 else split + row + 1
            cell = table.cell(row, col)
            cell_margins(cell, top=8, start=20, bottom=8, end=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(f"{number}. ")
            r.bold = True
            r.font.size = Pt(8)
            r = p.add_run(value)
            r.font.size = Pt(8)

def add_image(path, caption, width=8.35):
    assert path.exists(), path
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(str(path), width=Inches(width))
    c = doc.add_paragraph(caption)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(3)
    for r in c.runs:
        r.font.size = Pt(7.5)
        r.italic = True
        r.font.color.rgb = RGBColor(71, 85, 105)

def page(title, route, intro, image_path=None, caption="", steps=(), width=8.35):
    doc.add_page_break()
    doc.add_heading(title, 1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.add_run("Ruta  ").bold = True
    p.add_run(route)
    text(intro)
    if image_path:
        add_image(image_path, caption, width)
    add_legend(steps)

def cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)

def cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)

doc.add_paragraph("Manual de uso del sistema ETRAL por flujo", style="Title")
text("Guía para mantenedores, responsables de planta y trabajadores que registran la operación. El recorrido comienza con los datos maestros que alimentan el gemelo digital y continúa con la creación de órdenes, la ejecución y el análisis.")
text("En cada captura, los números naranjas indican el control que se debe pulsar o el campo cuyo texto se debe completar. Los datos mostrados son ejemplos; antes de guardar, confirme que correspondan al registro real.")
doc.add_heading("Orden recomendado de configuración", 1)
rows = [
    ("1", "Catálogos de inventario", "Categorías, unidades, marcas y familias"),
    ("2", "Materiales", "Descripción, stock, seguridad y ubicación"),
    ("3", "Productos y actividades", "Ruta de fabricación y tiempos estándar"),
    ("4", "Recursos", "Turnos, trabajadores, equipos y calendario"),
    ("5", "Clientes", "Datos comerciales que utilizará la orden"),
    ("6", "Orden CECO", "Cliente, producto y fechas del plan"),
    ("7", "Ejecución", "Asignación, avance, horas, incidencias y materiales"),
    ("8", "Análisis", "Indicadores y simulación del gemelo digital"),
]
table = doc.add_table(rows=1, cols=3)
table.autofit = False
table.columns[0].width, table.columns[1].width, table.columns[2].width = Inches(.6), Inches(2.5), Inches(6.6)
for i, label in enumerate(("Paso", "Interfaz", "Información que alimenta")):
    cell = table.rows[0].cells[i]
    cell.text = label
    cell_fill(cell, "1F2937")
    for r in cell.paragraphs[0].runs:
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.bold = True
for ridx, row in enumerate(rows):
    cells = table.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_margins(cells[i])
        if ridx % 2:
            cell_fill(cells[i], "F3F6F9")

page("1 Acceso al sistema", "Página de inicio de sesión", "Ingrese con la cuenta asignada. No comparta las credenciales.",
     ASSETS / "01_acceso_marcado.png", "1 Correo  2 Contraseña  3 Iniciar sesión",
     ["Escriba el correo asignado.", "Escriba la contraseña.", "Pulse Iniciar sesión."], 4.15)

page("2 Menú principal", "Inicio", "El menú lateral permite recorrer las áreas en el mismo orden del flujo.",
     WIDE / "01_inicio_marcado_amplio.png", "Opciones principales del sistema",
     ["Inicio: consulte el resumen y las alertas de la planta.", "Producción: gestione clientes, órdenes CECO y ejecución.", "Productos: mantenga plantillas, rutas y BOM.", "Fases y actividades: configure procesos y tareas.", "Inventario: administre materiales, catálogos y movimientos.", "Recursos: registre personal, turnos, equipos y calendario.", "Indicadores: revise resultados y exportaciones.", "Simulación: ejecute escenarios del gemelo digital."])

page("3 Catálogos de inventario", "Inventario > Catálogos", "Configure primero las listas que utilizarán materiales y productos.",
     FLOW / "02_catalogos_marcado.png", "Controles del mantenedor de catálogos",
     ["Pulse Inventario en el menú lateral.", "Abra la pestaña Catálogos.", "Elija Categorías, Unidades, Marcas o Familias.", "Escriba el nombre del nuevo registro.", "Pulse Agregar.", "Use Editar para corregir un registro existente."])

page("4 Registrar material", "Inventario > Materiales > Nuevo material", "Cree los materiales después de configurar sus catálogos.",
     FLOW / "03_nuevo_material_marcado.png", "Campos del registro de material",
     ["Escriba la descripción del material.", "Seleccione categoría y unidad de medida.", "Seleccione la marca.", "Ingrese stock inicial y stock de seguridad.", "Complete factor de servicio, variabilidad y plazo de reposición.", "Indique la ubicación.", "Pulse Guardar registro."])

page("5 Interfaz de productos", "Productos", "Esta pantalla contiene las plantillas maestras, rutas y listas de materiales.",
     FLOW / "04_productos.jpg", "Botones de importación, producto maestro, edición y BOM",
     ["Pulse Producto maestro para crear una plantilla.", "Seleccione una tarjeta para consultar su ruta y BOM.", "Use Editar plantilla o Material BOM según el dato que deba mantener."])

page("6 Registrar producto", "Productos > Producto maestro", "Defina la identidad del producto y las fases que recorrerá.",
     FLOW / "05_nuevo_producto_marcado.png", "Campos del producto y ruta de fabricación",
     ["Ingrese el código y seleccione la familia.", "Escriba el nombre del producto.", "Seleccione marca y unidad de salida.", "Ingrese los días objetivo.", "Marque las fases aplicables en la ruta de fabricación.", "Pulse Guardar registro."])

page("7 Interfaz de fases y actividades", "Fases y actividades", "Cada fase muestra su capacidad, tiempo estándar, actividades e inventario en proceso.",
     FLOW / "06_fases.jpg", "Fases, actividades, botones Editar y Añadir actividad",
     ["Pulse Añadir actividad para crear una tarea.", "Pulse Editar junto a una actividad para cambiar su nombre, tiempo o estado.", "Revise el control de calidad y el WIP de cada fase."])

page("8 Registrar actividad", "Fases y actividades > Añadir actividad", "Asocie la actividad con una fase y defina su tiempo estándar.",
     FLOW / "07_nueva_actividad_marcado.png", "Campos del registro de actividad",
     ["Seleccione la fase.", "Escriba el nombre de la actividad.", "Ingrese el tiempo estándar en minutos.", "Pulse Guardar registro."])

page("9 Pestañas de recursos", "Recursos", "Los recursos alimentan la capacidad disponible del gemelo digital.",
     FLOW / "08_recursos_tabs.jpg", "Pestañas Personal, Turnos, Equipos, Asignaciones, Calendario e Incidencias",
     ["Configure Turnos antes de crear trabajadores.", "Registre trabajadores y equipos; luego revise Asignaciones.", "Use Calendario para feriados o paradas e Incidencias para pérdidas de capacidad."])

page("10 Registrar turno", "Recursos > Turnos > Nuevo turno", "El turno determina el horario y el descanso que reducen o amplían la capacidad disponible.",
     FLOW / "10_nuevo_turno_marcado.png", "Campos del turno",
     ["Ingrese código y nombre.", "Defina hora de inicio y hora de fin.", "Ingrese los minutos de descanso.", "Pulse Guardar registro."])

page("11 Registrar trabajador", "Recursos > Personal > Nuevo trabajador", "El trabajador aporta horas y eficiencia a la capacidad laboral.",
     FLOW / "11_nuevo_trabajador_marcado.png", "Campos del trabajador",
     ["Ingrese código y estado.", "Escriba el nombre completo.", "Complete cargo y turno.", "Indique la especialidad.", "Revise eficiencia estimada y horas semanales.", "Pulse Guardar registro."])

page("12 Registrar equipo", "Recursos > Equipos > Nuevo equipo", "Los equipos se vinculan a una fase y aportan capacidad instalada.",
     FLOW / "13_nuevo_equipo_marcado.png", "Campos del equipo",
     ["Ingrese código y estado.", "Escriba el nombre del equipo.", "Seleccione la fase asociada.", "Indique capacidad semanal y próximo mantenimiento.", "Pulse Guardar registro."])

page("13 Registrar excepción de calendario", "Recursos > Calendario > Registrar excepción", "Registre feriados, paradas o jornadas reducidas que cambian la capacidad.",
     FLOW / "15_nueva_excepcion_marcado.png", "Campos de la excepción de calendario",
     ["Compruebe los días laborables base.", "Seleccione la fecha de excepción.", "Seleccione el tipo de excepción.", "Complete las horas si es jornada reducida.", "Agregue una descripción.", "Pulse Guardar registro."])

page("14 Interfaz de clientes", "Producción > Clientes", "El cliente se registra una vez y se selecciona después en cada orden CECO.",
     FLOW / "16_clientes.jpg", "Pestaña Clientes, buscador, Nuevo cliente y Editar",
     ["Pulse Clientes dentro de Producción.", "Use el buscador para evitar duplicados.", "Pulse Nuevo cliente o Editar."])

page("15 Registrar cliente", "Producción > Clientes > Nuevo cliente", "Complete los datos disponibles del cliente.",
     FLOW / "17_nuevo_cliente_marcado.png", "Campos del cliente",
     ["Ingrese razón social o nombre y documento.", "Complete la persona de contacto.", "Ingrese teléfono y correo.", "Pulse Guardar registro."])

page("16 Crear una orden CECO", "Producción > Nueva orden CECO", "La orden conecta un cliente con una plantilla de producto y las fechas del PMP.",
     FLOW / "18_nueva_orden_marcado.png", "Campos de la nueva orden",
     ["Confirme el correlativo CECO.", "Use Crear nuevo cliente si el cliente aún no existe.", "Seleccione el cliente.", "Use Crear nuevo producto si la plantilla aún no existe.", "Seleccione el producto.", "Ingrese Inicio PMP y Fecha pactada.", "Pulse Guardar registro."])

page("17 Vistas de producción", "Producción", "Las pestañas permiten consultar las órdenes por fase, cronograma, lista o cliente.",
     WIDE / "03_lista_marcada_amplio.png", "1 Lista  2 Buscador  3 Orden  4 Ver detalle",
     ["Abra la pestaña Lista.", "Busque por CECO, producto o cliente.", "Revise los datos y el estado de la orden.", "Pulse Ver detalle para abrir el pasaporte productivo."])

page("18 Detalle de la orden", "Producción > Lista > Ver detalle", "El pasaporte productivo concentra acciones, ruta, avance, materiales y calidad.",
    FLOW / "19_detalle_ceco_marcado.png", "1 Resumen de la orden  2 Botones de operación  3 Ruta del producto  4 Avance por etapa",
     ["Confirme cliente, fase, avance y fechas.", "Use Asignar trabajador, Registrar horas o Reportar incidencia.", "Revise la ruta completa del producto.", "Actualice el estado y avance de las actividades."], 8.35)

page("19 Asignar trabajador", "Detalle CECO > Asignar trabajador", "La asignación relaciona persona, orden, actividad, fecha y horas planificadas.",
     FLOW / "20_asignar_recurso_marcado.png", "Campos de la asignación",
     ["Seleccione el trabajador.", "Seleccione la orden CECO.", "Seleccione actividad y fecha.", "Ingrese horas planificadas y estado.", "Pulse Guardar registro."])

page("20 Actualizar actividades", "Detalle CECO > Actividades de la orden", "El avance por actividad determina el progreso de la fase y del CECO.",
     ASSETS / "08_avance_actividades_marcado.png", "1 Fase  2 Estado  3 Porcentaje  4 Completar",
     ["Abra la fase correcta.", "Seleccione el estado de la actividad.", "Ajuste el porcentaje de avance.", "Marque la casilla solo cuando la actividad esté terminada."], 4.7)

page("21 Registrar parte diario", "Detalle CECO > Registrar horas", "Registre el trabajo efectivamente realizado.",
     ASSETS / "06_registrar_horas_marcado.png", "Campos del parte diario",
     ["Confirme la orden CECO.", "Seleccione la fecha.", "Ingrese las horas trabajadas.", "Seleccione el responsable.", "Seleccione la actividad ejecutada.", "Pulse Guardar registro."], 4.7)

page("22 Reportar incidencia", "Detalle CECO > Reportar incidencia", "Registre los eventos que causan pérdida de capacidad o afectan la orden.",
     ASSETS / "07_reportar_incidencia_marcado.png", "Campos de la incidencia",
     ["Ingrese la fecha y hora.", "Seleccione la orden CECO si corresponde.", "Seleccione la fase afectada.", "Seleccione tipo y severidad.", "Seleccione el equipo si corresponde.", "Ingrese las horas de detención.", "Describa el evento y su impacto.", "Pulse Guardar registro."], 4.7)

page("23 Movimientos de inventario", "Inventario > Movimientos recientes", "El kardex conserva la trazabilidad de ingresos, ajustes, reservas, salidas y consumos.",
     WIDE / "09_movimientos.jpg", "Kardex y botón Registrar movimiento",
     ["Revise fecha, tipo, material, CECO, cantidad y detalle.", "El personal autorizado pulsa Registrar movimiento para un ingreso o ajuste.", "Las salidas de materiales reservados se ejecutan desde el detalle del CECO."])

page("24 Indicadores", "Indicadores", "El tablero presenta resultados para el período seleccionado y permite exportarlos.",
     WIDE / "12_indicadores.jpg", "Filtros, exportaciones e indicadores de producción",
     ["Seleccione Desde, Hasta y la serie.", "Revise unidades, horas, avance, PMP, cobertura y productividad.", "Confirme el período antes de exportar."])

page("25 Simulación del gemelo digital", "Simulación", "La simulación crea un escenario comparable sin modificar órdenes ni inventario real.",
     WIDE / "13_simulacion.jpg", "Parámetros del escenario",
     ["Defina horizonte, variabilidad, ausentismo, turnos y demanda.", "Ajuste materiales y prioridades CECO si el análisis lo requiere.", "Pulse Ejecutar simulación gemela."])

page("26 Resultados de simulación", "Simulación > Ejecutar simulación gemela", "Los resultados comparan la base con el escenario y muestran alertas accionables.",
     WIDE / "13_simulacion_resultados.jpg", "Resultados, alertas y pestañas de análisis",
     ["Compare órdenes terminables, cumplimiento PMP, lead time y quiebres.", "Revise capacidad, materiales, demanda, parámetros CECO, calibración y supuestos.", "Use el escenario para decidir; la ejecución no actualiza los registros reales."])

page("27 Comprobación final", "Antes de cerrar sesión", "Verifique que la información guardada refleje la operación real.",
     steps=["Mantenedores completos antes de crear órdenes.", "CECO, cliente, producto y fechas correctos.", "Asignaciones, avances, horas e incidencias registrados.", "Movimientos de almacén y liberaciones realizados por personal autorizado.", "Sesión cerrada en equipos compartidos."])

doc.save(OUT)
print(OUT)
